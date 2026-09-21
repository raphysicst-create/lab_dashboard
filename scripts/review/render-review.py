"""CPU preview fallback: read exported XLSX styles with openpyxl; draw with Pillow.

Artifact Tool's bundled Windows renderer terminates in Vulkan initialization on
this host. This helper never writes/changes workbook contents. It checks cell
dimensions, then previews every worksheet from the actual saved XLSX. It is not
a verification of Microsoft Excel's native visual rendering.
"""
from pathlib import Path
import json
import math
import os
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs' / 'review'
book = openpyxl.load_workbook(OUT / 'science_experiment_review.xlsx', data_only=False)
font_cache = {}

def find_fonts():
    """Use a Korean font from this machine, never another user's absolute path."""
    regular = os.environ.get('LAB_DASHBOARD_FONT')
    bold = os.environ.get('LAB_DASHBOARD_FONT_BOLD')
    if regular:
        paths = (Path(regular).expanduser(), Path(bold or regular).expanduser())
        if not all(path.is_file() for path in paths):
            raise FileNotFoundError('LAB_DASHBOARD_FONT / LAB_DASHBOARD_FONT_BOLD must point to existing font files.')
        return paths
    windows_fonts = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
    candidates = [
        (windows_fonts / 'malgun.ttf', windows_fonts / 'malgunbd.ttf'),
        (Path('/System/Library/Fonts/AppleSDGothicNeo.ttc'), Path('/System/Library/Fonts/AppleSDGothicNeo.ttc')),
        (Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'), Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc')),
        (Path('/usr/share/fonts/truetype/nanum/NanumGothic.ttf'), Path('/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf')),
    ]
    for regular_path, bold_path in candidates:
        if regular_path.is_file():
            return regular_path, bold_path if bold_path.is_file() else regular_path
    raise FileNotFoundError('Install a Korean font and set LAB_DASHBOARD_FONT to its .ttf/.otf/.ttc file path.')

FONT_PATHS = find_fonts()

def font_for(cell):
    size = round((cell.font.sz or 11) * 96 / 72)
    key = (size, bool(cell.font.b))
    if key not in font_cache:
        font_cache[key] = ImageFont.truetype(str(FONT_PATHS[1 if key[1] else 0]), size)
    return font_cache[key]

def color(value, default):
    return '#' + value.rgb[-6:] if value is not None and value.type == 'rgb' and isinstance(value.rgb, str) else default

def col_width(sheet, col):
    letter = get_column_letter(col)
    width = sheet.column_dimensions[letter].width
    return round((width or 13) * 7 + 5)

def row_height(sheet, row):
    return round((sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight or 15) * 96 / 72)

def text_lines(text, font, width, wrap):
    if not wrap:
        return str(text).split('\n')
    result = []
    for paragraph in str(text).split('\n'):
        current = ''
        for character in paragraph:
            if current and font.getlength(current + character) > width:
                result.append(current)
                current = character
            else:
                current += character
        result.append(current)
    return result

dimension_issues = []
for sheet in book.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            font = font_for(cell)
            lines = text_lines(cell.value, font, col_width(sheet, cell.column) - 12, cell.alignment.wrap_text)
            line_height = math.ceil(font.size * 1.25)
            needed_height = len(lines) * line_height + 8
            if needed_height > row_height(sheet, cell.row) + 2:
                dimension_issues.append({'sheet': sheet.title, 'cell': cell.coordinate, 'height': row_height(sheet, cell.row), 'needed': needed_height})

def render(name, address, filename):
    sheet = book[name]
    col1, row1, col2, row2 = range_boundaries(address)
    widths = [col_width(sheet, col) for col in range(col1, col2 + 1)]
    heights = [row_height(sheet, row) for row in range(row1, row2 + 1)]
    img = Image.new('RGB', (sum(widths) + 2, sum(heights) + 2), 'white')
    draw = ImageDraw.Draw(img)
    top = 1
    for row, height in zip(range(row1, row2 + 1), heights):
        left = 1
        for col, width in zip(range(col1, col2 + 1), widths):
            cell = sheet.cell(row, col)
            fill = color(cell.fill.fgColor, '#FFFFFF') if cell.fill.patternType == 'solid' else '#FFFFFF'
            draw.rectangle((left, top, left + width, top + height), fill=fill)
            left += width
        left = 1
        for col, width in zip(range(col1, col2 + 1), widths):
            cell = sheet.cell(row, col)
            if cell.value is not None:
                font = font_for(cell)
                text_color = color(cell.font.color, '#172B40')
                available_width = width - 12
                if not cell.alignment.wrap_text and isinstance(cell.value, str):
                    for next_col in range(col + 1, col2 + 1):
                        if sheet.cell(row, next_col).value is not None:
                            break
                        available_width += col_width(sheet, next_col)
                lines = text_lines(cell.value, font, available_width, cell.alignment.wrap_text)
                line_height = math.ceil(font.size * 1.25)
                vertical = cell.alignment.vertical
                y = top + (height - len(lines) * line_height) / 2 if vertical == 'center' else top + 4
                for line in lines:
                    x = left + 6
                    if cell.alignment.horizontal == 'center':
                        x = left + (width - font.getlength(line)) / 2
                    elif cell.alignment.horizontal == 'right':
                        x = left + width - font.getlength(line) - 6
                    draw.text((x, y), line, font=font, fill=text_color)
                    y += line_height
            left += width
        top += height
    img.save(OUT / filename)

for name, address, filename in [
    ('탐구활동 검토', 'A2:H9', 'review-original.png'),
    ('탐구활동 검토', 'I5:N9', 'review-inputs.png'),
    ('탐구활동 검토', 'O5:W7', 'review-corrections.png'),
    ('검토 방법', 'A2:B16', 'review-guide.png'),
    ('검토 방법', 'A17:B27', 'review-guide-continuation.png'),
]:
    render(name, address, filename)
report = {'renderer': 'Pillow CPU from saved XLSX through read-only openpyxl', 'sheets_previewed': book.sheetnames, 'dimension_issues': dimension_issues, 'native_excel_visual_check': 'not performed'}
(OUT / 'render-validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'previews': 5, 'dimension_issues': len(dimension_issues), 'first_issues': dimension_issues[:10]}, ensure_ascii=True))
